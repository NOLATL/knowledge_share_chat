"""Document upload parsing, chunking, embeddings, and document-derived extraction."""

from __future__ import annotations

import csv
import os
import uuid
from pathlib import Path

import pandas as pd

from api import answers_repo, document_repo, openai_client, repo
from playbook_shared.analytics.openai_ops import embed_texts, embedding_model
from playbook_shared.analytics.models import EMBEDDING_DIM

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".pptx", ".xlsx", ".csv"}
DEFAULT_MAX_UPLOAD_MB = 10
CHUNK_WORDS = 500
CHUNK_OVERLAP_WORDS = 75

_DOCUMENT_EXTRACTION_INSTRUCTIONS = (
    "You extract structured facts from uploaded finance/accounting discovery documents. "
    "Return values ONLY for slots supported by the list below. Do not invent values. "
    "Use concise values grounded in the document text and set confidence in [0,1]."
)


class DocumentError(ValueError):
    """Raised for user-correctable document upload errors."""


def max_upload_bytes() -> int:
    raw = os.environ.get("DOCUMENT_MAX_UPLOAD_MB")
    if raw is None:
        mb = DEFAULT_MAX_UPLOAD_MB
    else:
        try:
            mb = int(raw)
        except ValueError as exc:
            raise ValueError("DOCUMENT_MAX_UPLOAD_MB must be an integer") from exc
    if mb <= 0:
        raise ValueError("DOCUMENT_MAX_UPLOAD_MB must be greater than 0")
    return mb * 1024 * 1024


def upload_root() -> Path:
    value = os.environ.get("UPLOAD_VOLUME_PATH")
    if not value:
        raise RuntimeError("UPLOAD_VOLUME_PATH is not set")
    return Path(value)


def safe_filename(filename: str) -> str:
    name = Path(filename).name.strip()
    if not name:
        raise DocumentError("filename is required")
    return "".join(ch if ch.isalnum() or ch in "._- " else "_" for ch in name)


def infer_doc_type(filename: str, explicit: str | None = None) -> str:
    if explicit:
        return explicit
    suffix = Path(filename).suffix.lower()
    if suffix in {".xlsx", ".csv"}:
        return "financials"
    if suffix == ".pptx":
        return "deck"
    if suffix in {".pdf", ".docx"}:
        return "notes"
    return "unknown"


def validate_extension(filename: str) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise DocumentError(f"unsupported file type: {suffix or 'none'}")
    return suffix


def parse_file(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return _parse_csv(path)
    if suffix == ".xlsx":
        return _parse_xlsx(path)
    if suffix in {".pdf", ".docx", ".pptx"}:
        return _parse_docling(path)
    raise DocumentError(f"unsupported file type: {suffix or 'none'}")


def _parse_csv(path: Path) -> str:
    try:
        df = pd.read_csv(path)
    except UnicodeDecodeError:
        df = pd.read_csv(path, encoding="latin-1")
    return _dataframe_to_markdown(df, title=path.name)


def _parse_xlsx(path: Path) -> str:
    sheets = pd.read_excel(path, sheet_name=None)
    parts = []
    for name, df in sheets.items():
        parts.append(_dataframe_to_markdown(df, title=f"Sheet: {name}"))
    return "\n\n".join(parts)


def _parse_docling(path: Path) -> str:
    try:
        from docling.document_converter import DocumentConverter

        result = DocumentConverter().convert(path)
        document = result.document
        if hasattr(document, "export_to_markdown"):
            return str(document.export_to_markdown())
        if hasattr(document, "export_to_text"):
            return str(document.export_to_text())
        return str(document)
    except ModuleNotFoundError:
        return _parse_without_docling_optional_backends(path)


def _parse_without_docling_optional_backends(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _parse_pdf_direct(path)
    if suffix == ".docx":
        return _parse_docx_direct(path)
    if suffix == ".pptx":
        return _parse_pptx_direct(path)
    raise DocumentError(f"unsupported file type: {suffix or 'none'}")


def _parse_pdf_direct(path: Path) -> str:
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(path)
    pages = []
    for index, page in enumerate(pdf, start=1):
        textpage = page.get_textpage()
        text = textpage.get_text_range()
        if text.strip():
            pages.append(f"## Page {index}\n{text.strip()}")
    return "\n\n".join(pages)


def _parse_docx_direct(path: Path) -> str:
    from docx import Document

    doc = Document(path)
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table_index, table in enumerate(doc.tables, start=1):
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if rows:
            parts.append(f"## Table {table_index}")
            parts.append(_markdown_row(rows[0]))
            parts.append(_markdown_row(["---"] * len(rows[0])))
            parts.extend(_markdown_row(row) for row in rows[1:])
    return "\n\n".join(parts)


def _parse_pptx_direct(path: Path) -> str:
    from pptx import Presentation

    deck = Presentation(path)
    slides = []
    for index, slide in enumerate(deck.slides, start=1):
        texts = []
        for shape in slide.shapes:
            text = getattr(shape, "text", "")
            if text.strip():
                texts.append(text.strip())
        if texts:
            slides.append(f"## Slide {index}\n" + "\n".join(texts))
    return "\n\n".join(slides)


def _dataframe_to_markdown(df: pd.DataFrame, *, title: str) -> str:
    clean = df.fillna("")
    headers = [str(c) for c in clean.columns]
    rows = clean.astype(str).values.tolist()
    output = [f"## {title}", _markdown_row(headers), _markdown_row(["---"] * len(headers))]
    output.extend(_markdown_row(row) for row in rows)
    return "\n".join(output)


def _markdown_row(values: list[str]) -> str:
    escaped = [v.replace("|", "\\|").replace("\n", " ") for v in values]
    return "| " + " | ".join(escaped) + " |"


def chunk_text(text: str, *, chunk_words: int = CHUNK_WORDS, overlap_words: int = CHUNK_OVERLAP_WORDS) -> list[str]:
    words = text.split()
    if not words:
        return []
    chunks: list[str] = []
    start = 0
    step = max(1, chunk_words - overlap_words)
    while start < len(words):
        chunks.append(" ".join(words[start : start + chunk_words]))
        start += step
    return chunks


def write_chunks(
    *,
    document_id: uuid.UUID,
    company_id: uuid.UUID,
    session_id: uuid.UUID | str,
    doc_type: str,
    parsed_text: str,
) -> int:
    chunks = chunk_text(parsed_text)
    document_repo.delete_document_chunks(document_id)
    if not chunks:
        return 0

    vectors = embed_texts(chunks)
    if any(len(vector) != EMBEDDING_DIM for vector in vectors):
        raise ValueError(f"document embedding dimension must be {EMBEDDING_DIM}")

    for index, (chunk, vector) in enumerate(zip(chunks, vectors, strict=True)):
        document_repo.insert_document_chunk(
            document_id=document_id,
            chunk_text=chunk,
            embedding=vector,
            metadata={
                "company_id": str(company_id),
                "session_id": str(session_id),
                "document_id": str(document_id),
                "doc_type": doc_type,
                "chunk_index": index,
                "embedding_model": embedding_model(),
            },
        )
    return len(chunks)


async def extract_document_answers(session_id: uuid.UUID | str, parsed_text: str) -> list[str]:
    questions = repo.load_active_questions()
    active_keys = {q["slot_key"] for q in questions}
    slot_list = "\n".join(f"- {q['slot_key']}: {q['text']}" for q in questions)
    messages = [
        {"role": "system", "content": f"{_DOCUMENT_EXTRACTION_INSTRUCTIONS}\n\nSlots:\n{slot_list}"},
        {"role": "user", "content": parsed_text[:12000]},
    ]
    result = await openai_client.parse_extraction(messages)

    saved: list[str] = []
    for ex in result.extractions:
        if ex.slot_key not in active_keys:
            continue
        answers_repo.save_answer(
            session_id,
            ex.slot_key,
            parsed_text[:4000],
            {"value": ex.value},
            ex.confidence,
            "document",
        )
        saved.append(ex.slot_key)
    return saved


def read_limited_upload(data: bytes) -> bytes:
    limit = max_upload_bytes()
    if len(data) > limit:
        raise DocumentError(f"file exceeds {limit // (1024 * 1024)} MB limit")
    return data


def sniff_csv_text(path: Path) -> bool:
    try:
        with path.open(newline="", encoding="utf-8") as handle:
            csv.Sniffer().sniff(handle.read(2048))
        return True
    except Exception:
        return False
