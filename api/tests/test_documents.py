from __future__ import annotations

import uuid

import psycopg
from fastapi.testclient import TestClient

from api import document_processing, document_rag, openai_client
from api.main import app
from playbook_shared.extraction_schemas import ExtractedSlot, TurnExtraction

client = TestClient(app)


def _session_id() -> str:
    return client.post("/sessions").json()["id"]


def _vector(value: float = 0.01) -> list[float]:
    return [value] * 1536


async def _no_document_answers(session_id, parsed_text):
    return []


def test_upload_rejects_unknown_session(chat_db, tmp_path, monkeypatch):
    monkeypatch.setenv("UPLOAD_VOLUME_PATH", str(tmp_path))
    resp = client.post(
        f"/sessions/{uuid.uuid4()}/documents",
        files={"file": ("sample.csv", b"a,b\n1,2\n", "text/csv")},
    )
    assert resp.status_code == 404


def test_upload_rejects_unsupported_extension(chat_db, tmp_path, monkeypatch):
    monkeypatch.setenv("UPLOAD_VOLUME_PATH", str(tmp_path))
    session_id = _session_id()
    resp = client.post(
        f"/sessions/{session_id}/documents",
        files={"file": ("diagram.png", b"x", "image/png")},
    )
    assert resp.status_code == 400
    assert "unsupported" in resp.json()["detail"]


def test_upload_rejects_over_limit_file(chat_db, tmp_path, monkeypatch):
    monkeypatch.setenv("UPLOAD_VOLUME_PATH", str(tmp_path))
    monkeypatch.setenv("DOCUMENT_MAX_UPLOAD_MB", "1")
    session_id = _session_id()
    resp = client.post(
        f"/sessions/{session_id}/documents",
        files={"file": ("large.csv", b"x" * (1024 * 1024 + 1), "text/csv")},
    )
    assert resp.status_code == 400
    assert "limit" in resp.json()["detail"]


def test_csv_upload_parses_writes_file_and_chunks(chat_db, tmp_path, monkeypatch):
    monkeypatch.setenv("UPLOAD_VOLUME_PATH", str(tmp_path))
    monkeypatch.setattr(document_processing, "embed_texts", lambda chunks: [_vector() for _ in chunks])
    monkeypatch.setattr(document_processing, "extract_document_answers", _no_document_answers)

    session_id = _session_id()
    resp = client.post(
        f"/sessions/{session_id}/documents",
        files={"file": ("budget.csv", b"metric,value\nClose days,12\nManual hours,80\n", "text/csv")},
    )

    assert resp.status_code == 200
    body = resp.json()
    assert body["status"] == "parsed"
    assert body["chunk_count"] >= 1
    assert (tmp_path / session_id / body["id"] / "budget.csv").is_file()

    with psycopg.connect(chat_db) as conn, conn.cursor() as cur:
        cur.execute("SELECT parsed_text FROM documents WHERE id = %s", (body["id"],))
        parsed = cur.fetchone()[0]
    assert "| metric | value |" in parsed
    assert "| Close days | 12 |" in parsed


def test_parser_failure_sets_error_status(chat_db, tmp_path, monkeypatch):
    monkeypatch.setenv("UPLOAD_VOLUME_PATH", str(tmp_path))
    monkeypatch.setattr(document_processing, "parse_file", lambda path: (_ for _ in ()).throw(ValueError("bad file")))

    session_id = _session_id()
    resp = client.post(
        f"/sessions/{session_id}/documents",
        files={"file": ("broken.csv", b"metric,value\n", "text/csv")},
    )

    assert resp.status_code == 422
    docs = client.get(f"/sessions/{session_id}/documents").json()
    assert docs[0]["status"] == "error"


def test_document_extraction_writes_document_source_answer(chat_db, tmp_path, monkeypatch):
    monkeypatch.setenv("UPLOAD_VOLUME_PATH", str(tmp_path))
    monkeypatch.setattr(document_processing, "embed_texts", lambda chunks: [_vector() for _ in chunks])

    async def fake_parse(messages):
        return TurnExtraction(
            extractions=[ExtractedSlot(slot_key="close.days", value="12 days", confidence=0.88)]
        )

    monkeypatch.setattr(openai_client, "parse_extraction", fake_parse)
    session_id = _session_id()
    resp = client.post(
        f"/sessions/{session_id}/documents",
        files={"file": ("close.csv", b"metric,value\nClose days,12\n", "text/csv")},
    )

    assert resp.status_code == 200
    with psycopg.connect(chat_db) as conn, conn.cursor() as cur:
        cur.execute(
            "SELECT structured_json, source FROM answers WHERE session_id = %s AND slot_key = 'close.days'",
            (session_id,),
        )
        row = cur.fetchone()
    assert row[0] == {"value": "12 days"}
    assert row[1] == "document"


def test_retrieval_is_company_scoped(chat_db, tmp_path, monkeypatch):
    monkeypatch.setenv("UPLOAD_VOLUME_PATH", str(tmp_path))
    monkeypatch.setattr(document_processing, "embed_texts", lambda chunks: [_vector() for _ in chunks])
    monkeypatch.setattr(document_processing, "extract_document_answers", _no_document_answers)
    monkeypatch.setattr(document_rag, "embed_texts", lambda texts: [_vector()])

    session_a = _session_id()
    session_b = _session_id()
    client.post(
        f"/sessions/{session_a}/documents",
        files={"file": ("a.csv", b"metric,value\nClose days,12\n", "text/csv")},
    )
    client.post(
        f"/sessions/{session_b}/documents",
        files={"file": ("b.csv", b"metric,value\nClose days,4\n", "text/csv")},
    )

    context = document_rag.retrieve_document_context(session_a, "close days")
    assert "a.csv" in context
    assert "b.csv" not in context


def test_chat_prompt_includes_retrieved_document_fact(chat_db, tmp_path, monkeypatch):
    monkeypatch.setenv("UPLOAD_VOLUME_PATH", str(tmp_path))
    monkeypatch.setattr(document_processing, "embed_texts", lambda chunks: [_vector() for _ in chunks])
    monkeypatch.setattr(document_processing, "extract_document_answers", _no_document_answers)
    monkeypatch.setattr(document_rag, "embed_texts", lambda texts: [_vector()])

    async def fake_parse(messages):
        return TurnExtraction(extractions=[])

    captured: dict = {}

    async def fake_stream(messages):
        captured["messages"] = messages
        yield "ok"

    monkeypatch.setattr(openai_client, "parse_extraction", fake_parse)
    monkeypatch.setattr(openai_client, "stream_chat", fake_stream)

    session_id = _session_id()
    client.post(
        f"/sessions/{session_id}/documents",
        files={"file": ("close-calendar.csv", b"metric,value\nClose days,12\n", "text/csv")},
    )
    resp = client.post(f"/sessions/{session_id}/messages", json={"message": "What do you know about close?"})

    assert resp.status_code == 200
    system = captured["messages"][0]["content"]
    assert "Uploaded document context" in system
    assert "close-calendar.csv" in system
    assert "Close days" in system


def test_docx_and_pptx_fallback_parsers_preserve_text(tmp_path):
    from docx import Document
    from pptx import Presentation

    docx_path = tmp_path / "notes.docx"
    doc = Document()
    doc.add_paragraph("Close takes 12 days.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Manual hours"
    table.cell(1, 1).text = "80"
    doc.save(docx_path)

    pptx_path = tmp_path / "deck.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Finance close"
    deck.save(pptx_path)

    docx_text = document_processing._parse_without_docling_optional_backends(docx_path)
    pptx_text = document_processing._parse_without_docling_optional_backends(pptx_path)

    assert "Close takes 12 days." in docx_text
    assert "| Metric | Value |" in docx_text
    assert "Finance close" in pptx_text
